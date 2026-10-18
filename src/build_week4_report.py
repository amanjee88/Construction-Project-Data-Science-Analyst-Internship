"""Build the Week 4 editable Word report from generated Week 4 outputs."""
from __future__ import annotations
import json
from pathlib import Path
import pandas as pd
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "reports" / "Week_4_Data_Visualization_Reporting_and_Actionable_Insights.docx"
PROCESSED = ROOT / "data" / "processed"
FIG = ROOT / "reports" / "figures"

NAVY = "17324D"
BLUE = "2E75B6"
TEAL = "2A9D8F"
LIGHT = "EAF0F6"
MID = "D7E2EC"
GREY = "5B6573"
ORANGE = "E58B2A"


def set_cell_shading(cell, fill: str):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tcPr.append(shd)


def set_cell_margins(cell, top=90, start=100, bottom=90, end=100):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    tcMar = tcPr.find(qn("w:tcMar"))
    if tcMar is None:
        tcMar = OxmlElement("w:tcMar")
        tcPr.append(tcMar)
    for m, v in [("top", top), ("start", start), ("bottom", bottom), ("end", end)]:
        node = tcMar.find(qn(f"w:{m}"))
        if node is None:
            node = OxmlElement(f"w:{m}")
            tcMar.append(node)
        node.set(qn("w:w"), str(v)); node.set(qn("w:type"), "dxa")


def keep_table_header(row):
    trPr = row._tr.get_or_add_trPr()
    repeat = OxmlElement("w:tblHeader")
    repeat.set(qn("w:val"), "true")
    trPr.append(repeat)
    cant_split = OxmlElement("w:cantSplit")
    trPr.append(cant_split)


def add_page_number(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run("Page ")
    fldChar1 = OxmlElement("w:fldChar"); fldChar1.set(qn("w:fldCharType"), "begin")
    instrText = OxmlElement("w:instrText"); instrText.set(qn("xml:space"), "preserve"); instrText.text = " PAGE "
    fldChar2 = OxmlElement("w:fldChar"); fldChar2.set(qn("w:fldCharType"), "end")
    run._r.append(fldChar1); run._r.append(instrText); run._r.append(fldChar2)


def add_table(doc, headers, rows, widths=None, font_size=8.5):
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    hdr = table.rows[0]; keep_table_header(hdr)
    for i, h in enumerate(headers):
        c = hdr.cells[i]; c.text = str(h); set_cell_shading(c, NAVY); set_cell_margins(c)
        c.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        for p in c.paragraphs:
            p.paragraph_format.space_after = Pt(0)
            for r in p.runs:
                r.font.name = "Aptos"; r.font.size = Pt(font_size); r.font.bold = True; r.font.color.rgb = RGBColor(255,255,255)
        if widths: c.width = Inches(widths[i])
    for ridx, row in enumerate(rows):
        cells = table.add_row().cells
        trPr = table.rows[-1]._tr.get_or_add_trPr()
        trPr.append(OxmlElement("w:cantSplit"))
        for i, value in enumerate(row):
            cells[i].text = str(value)
            cells[i].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_margins(cells[i])
            if ridx % 2 == 1: set_cell_shading(cells[i], "F4F7FA")
            if widths: cells[i].width = Inches(widths[i])
            for p in cells[i].paragraphs:
                p.paragraph_format.space_after = Pt(0)
                p.paragraph_format.space_before = Pt(0)
                for r in p.runs:
                    r.font.name = "Aptos"; r.font.size = Pt(font_size); r.font.color.rgb = RGBColor(40, 50, 60)
    doc.add_paragraph().paragraph_format.space_after = Pt(1)
    return table


def add_bullet(doc, text, level=0):
    p = doc.add_paragraph(style="List Bullet" if level == 0 else "List Bullet 2")
    p.paragraph_format.space_after = Pt(3)
    p.add_run(text)
    return p


def add_number(doc, text):
    p = doc.add_paragraph(style="List Number")
    p.paragraph_format.space_after = Pt(3)
    p.add_run(text)
    return p


def add_callout(doc, title, text, fill=LIGHT):
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = table.cell(0,0); set_cell_shading(cell, fill); set_cell_margins(cell, top=140, start=180, bottom=140, end=180)
    p = cell.paragraphs[0]; p.paragraph_format.space_after = Pt(3)
    r = p.add_run(title); r.bold = True; r.font.color.rgb = RGBColor.from_string(NAVY); r.font.size = Pt(10)
    p2 = cell.add_paragraph(text); p2.paragraph_format.space_after = Pt(0)
    for r in p2.runs: r.font.size = Pt(9); r.font.color.rgb = RGBColor.from_string(GREY)
    doc.add_paragraph().paragraph_format.space_after = Pt(1)


def add_figure(doc, filename, caption, width=6.35):
    path = FIG / filename
    if not path.exists():
        raise FileNotFoundError(path)
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(2)
    p.add_run().add_picture(str(path), width=Inches(width))
    cp = doc.add_paragraph(); cp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cp.paragraph_format.space_after = Pt(6)
    run = cp.add_run(caption); run.italic = True; run.font.size = Pt(8); run.font.color.rgb = RGBColor.from_string(GREY)


def add_heading(doc, text, level=1):
    p = doc.add_heading(text, level=level)
    p.paragraph_format.keep_with_next = True
    major = text == "Executive Summary" or text.startswith(("1.", "2.", "3.", "4.", "6.", "7."))
    if level == 1 and major:
        p.paragraph_format.page_break_before = True
    return p


def add_note(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(5)
    r = p.add_run(text); r.italic = True; r.font.size = Pt(8.5); r.font.color.rgb = RGBColor.from_string(GREY)


def main():
    summary = json.loads((PROCESSED / "week4_executive_summary.json").read_text(encoding="utf-8"))
    w1 = summary["week1_data_acquisition_sample"]
    w2 = summary["week2_eda_sample"]
    w3 = summary["week3_risk_sample"]
    scenarios = pd.read_csv(PROCESSED / "week3_scenario_analysis.csv")
    project_types = pd.read_csv(PROCESSED / "week4_project_type_summary.csv")
    risks = pd.read_csv(PROCESSED / "week4_risk_priority_register.csv")
    recs = pd.read_csv(PROCESSED / "week4_recommendation_tracker.csv")

    doc = Document()
    sec = doc.sections[0]
    sec.top_margin = Inches(0.68); sec.bottom_margin = Inches(0.68)
    sec.left_margin = Inches(0.78); sec.right_margin = Inches(0.78)
    sec.header_distance = Inches(0.32); sec.footer_distance = Inches(0.32)

    styles = doc.styles
    styles["Normal"].font.name = "Aptos"; styles["Normal"]._element.rPr.rFonts.set(qn("w:eastAsia"), "Aptos")
    styles["Normal"].font.size = Pt(9.5); styles["Normal"].font.color.rgb = RGBColor(42, 50, 60)
    styles["Normal"].paragraph_format.space_after = Pt(5); styles["Normal"].paragraph_format.line_spacing = 1.08
    for sname, size, color in [("Title", 29, NAVY), ("Heading 1", 17, NAVY), ("Heading 2", 12, BLUE), ("Heading 3", 10, TEAL)]:
        st = styles[sname]; st.font.name = "Aptos Display"; st._element.rPr.rFonts.set(qn("w:eastAsia"), "Aptos Display")
        st.font.size = Pt(size); st.font.bold = True; st.font.color.rgb = RGBColor.from_string(color)
        st.paragraph_format.space_before = Pt(10 if sname != "Title" else 0); st.paragraph_format.space_after = Pt(5)
    # Running header/footer
    header = sec.header.paragraphs[0]
    header.text = "CONSTRUCTION PROJECT DATA SCIENCE  |  WEEK 4 FINAL REPORT"
    header.alignment = WD_ALIGN_PARAGRAPH.LEFT
    for r in header.runs: r.font.name = "Aptos"; r.font.size = Pt(8); r.font.bold = True; r.font.color.rgb = RGBColor.from_string(BLUE)
    footer = sec.footer.paragraphs[0]
    footer.add_run("Synthetic-data portfolio exercise  |  ")
    for r in footer.runs: r.font.size = Pt(8); r.font.color.rgb = RGBColor.from_string(GREY)
    add_page_number(footer)

    # Cover
    p = doc.add_paragraph(); p.paragraph_format.space_before = Pt(68); p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    r = p.add_run("CONSTRUCTION\nPROJECT ANALYTICS")
    r.bold = True; r.font.name = "Aptos Display"; r.font.size = Pt(14); r.font.color.rgb = RGBColor.from_string(TEAL)
    p = doc.add_paragraph(style="Title")
    p.paragraph_format.space_before = Pt(12)
    p.add_run("Data Visualization, Reporting, and Actionable Insights")
    p2 = doc.add_paragraph(); p2.paragraph_format.space_before = Pt(8); p2.paragraph_format.space_after = Pt(17)
    r = p2.add_run("Week 4 | Stakeholder Decision Report")
    r.font.size = Pt(16); r.font.color.rgb = RGBColor.from_string(BLUE)
    p3 = doc.add_paragraph("Synthesis of data acquisition, exploratory analysis, predictive modeling, and risk response planning")
    p3.paragraph_format.space_after = Pt(20)
    for r in p3.runs: r.font.size = Pt(12); r.font.color.rgb = RGBColor.from_string(GREY)
    add_callout(doc, "Purpose", "Translate the Week 1-3 analytical work into a clear performance narrative, communicate the key patterns visually, and propose practical project-controls actions.", fill="EAF4F5")
    doc.add_paragraph("Prepared for: Construction project stakeholders, project controls teams, site managers, procurement, cost control, and HSE leads")
    doc.add_paragraph("Scope: Python-based analysis using explicitly simulated project data; all results are illustrative and not actual industry observations.")
    doc.add_paragraph("Prepared as the final deliverable for a virtual internship task.")

    # Executive summary
    add_heading(doc, "Executive Summary", 1)
    doc.add_paragraph(
        "This report brings together the four-week construction analytics workflow: a data acquisition and quality strategy (Week 1), exploratory data analysis and initial statistical modeling (Week 2), and predictive risk modeling (Week 3). The final emphasis is not a claim that the models are deployment-ready; it is a stakeholder-ready demonstration of how data can be organized, interpreted, validated, and converted into management action."
    )
    doc.add_paragraph(
        f"In the Week 2 synthetic EDA sample of {w2['rows']} project records, mean schedule delay was {w2['mean_schedule_delay_days']:.1f} days, the on-time indicator rate was {w2['on_time_rate_pct']:.1f}%, and mean cost overrun was {w2['mean_cost_overrun_pct']:.2f}%. Procurement delay and schedule delay had a positive Pearson correlation of {w2['correlation_procurement_schedule']:.3f}; material cost index and cost overrun had a correlation of {w2['correlation_material_cost_overrun']:.3f}. These are associations in simulated data and do not establish causality."
    )
    doc.add_paragraph(
        f"The separate Week 3 synthetic risk sample of {w3['rows']} records was used to compare delay regressors and cost-overrun classifiers. Linear Regression produced a holdout delay MAE of {w3['delay_linear_regression']['holdout_MAE_days']:.2f} days and R-squared of {w3['delay_linear_regression']['holdout_R2']:.3f}. Logistic Regression reached {w3['cost_risk_logistic_regression']['recall']:.3f} recall and {w3['cost_risk_logistic_regression']['roc_auc']:.3f} ROC-AUC for the 10% cost-overrun flag in that sample. The metrics are workflow demonstrations, not promised operational performance."
    )
    add_figure(doc, "week4_01_executive_kpi_dashboard.png", "Figure 1. Performance snapshot derived from the Week 2 simulated EDA sample.")
    add_callout(doc, "Executive takeaway", "Build a weekly, owner-led project-controls rhythm around procurement readiness, schedule variance, forecast final cost, design changes/rework, workforce availability, site conditions, and safety. Use predictive scores as prompts for investigation—not as substitutes for engineering judgement, contractual controls, or safety procedures.")
    # Evidence & methodology follows naturally after the executive takeaway.
    add_heading(doc, "1. Scope, Evidence Base, and Data Limitations", 1)
    doc.add_paragraph("The report synthesizes the objectives, methods, and themes from the earlier tasks. The Week 1 starter data, Week 2 EDA data, and Week 3 risk-modeling data are separate synthetic fixtures with different sample sizes, fields, and generation assumptions. They are not the same projects followed longitudinally and have not been row-merged. Comparisons are therefore descriptive within each sample, not comparisons of real project populations across weeks.")
    add_table(doc, ["Workstream", "Sample / structure", "Use in final report", "Important limitation"], [
        ["Week 1: acquisition and preparation", f"{w1['rows']} rows; {w1['columns']} columns", "Data dictionary, staged-data concept, quality review", f"{w1['exact_duplicate_rows']} exact duplicate row(s); {w1['rows_with_duplicate_project_id']} rows tied to duplicate project IDs; {w1['rows_flagged_for_manual_review']} rows flagged for manual review."],
        ["Week 2: EDA and statistical modeling", f"{w2['rows']} rows; 19 columns", "Performance indicators, project-type comparison, correlations, cost-budget comparison", "Synthetic sample generated for analysis practice; no observed industry prevalence can be inferred."],
        ["Week 3: predictive risk modeling", f"{w3['rows']} rows; 20 columns", "Holdout and five-fold CV metrics; risk classification; what-if scenario results", "Separate synthetic sample; designed patterns can make metrics unlike real projects."],
    ], widths=[1.25, 1.2, 1.75, 2.25], font_size=8)
    add_heading(doc, "Data-quality principles carried forward", 2)
    for b in [
        "Preserve raw source extracts and record each cleaning rule; never silently overwrite the original data.",
        "Check duplicate project IDs, missing values, units, date ordering, valid ranges, and the distinction between planned, forecast, and actual fields.",
        "Treat anomalous records as review candidates rather than automatically deleting them; confirm with a project data owner.",
        "Only train models on fields that would be known at the forecast date. Exclude future information that would leak the outcome into the model."
    ]: add_bullet(doc, b)
    add_heading(doc, "Methods and Python tools", 2)
    add_table(doc, ["Tool / method", "Role in the workflow", "Stakeholder translation"], [
        ["pandas / NumPy", "Read, validate, group, summarize, and calculate indicators", "Turn tabular records into consistent, auditable summaries."],
        ["Matplotlib", "Create charts and the executive KPI view", "Show patterns and exceptions in a form suitable for meetings."],
        ["scikit-learn pipelines", "Imputation/encoding where needed, train-test split, five-fold validation, model metrics", "Evaluate model performance on records not used for fitting and compare stability."],
        ["Linear Regression", "Baseline estimate of schedule-delay days", "Quantify how inputs relate to a continuous outcome under a linear approximation."],
        ["Logistic Regression / Random Forest", "Compare cost-overrun risk probabilities and classification trade-offs", "Prioritize review, while accounting for missed risks and false alarms."],
    ], widths=[1.25, 2.55, 2.65], font_size=8)

    # Findings with visuals
    add_heading(doc, "2. Finding One: Schedule Performance Needs Active Control", 1)
    doc.add_paragraph(
        f"The Week 2 sample had a mean schedule delay of {w2['mean_schedule_delay_days']:.2f} days and a median of {w2['median_schedule_delay_days']:.1f} days. The on-time field was true for {w2['on_time_rate_pct']:.1f}% of records. This sample-specific pattern supports making schedule readiness and variance review a core dashboard function, but the percentages must not be presented as a construction-industry benchmark."
    )
    add_figure(doc, "week4_02_delay_by_project_type.png", "Figure 2. Mean schedule delay by project type in the Week 2 synthetic EDA sample.")
    top = project_types.iloc[0]; bottom = project_types.iloc[-1]
    doc.add_paragraph(f"Within this sample, {top['project_type']} projects had the highest mean delay ({top['mean_schedule_delay_days']:.1f} days), while {bottom['project_type']} projects had the lowest ({bottom['mean_schedule_delay_days']:.1f} days). Differences between categories can reflect sample construction and project complexity; managers should use such views to ask which scope, procurement, weather, approval, or resource conditions need review—not to assume that a project category inherently causes delay.")
    add_heading(doc, "Operational response", 2)
    add_bullet(doc, "Use a phase-level look-ahead and identify the next four to six weeks of critical activities, their constraints, responsible owners, and recovery options.")
    add_bullet(doc, "Review schedule variance together with open RFIs, approvals, procurement promises, workforce coverage, weather windows, and inspection/rework events.")
    add_bullet(doc, "Record the reason for every material forecast change so actual results can later be used to validate and improve the forecasting process.")

    add_heading(doc, "3. Finding Two: Procurement and Cost Signals Are Useful Early Warnings", 1)
    doc.add_paragraph(
        f"The Pearson correlation between procurement delay and schedule delay was {w2['correlation_procurement_schedule']:.3f} in the Week 2 sample. The relationship is positive: records with larger procurement delays tended, on average, to show larger schedule delays. Material cost index and cost overrun had a correlation of {w2['correlation_material_cost_overrun']:.3f}. Both results can guide questions for a project review, but correlation does not prove that one variable alone produced the outcome; project scale, design maturity, contract terms, and other unobserved factors may contribute."
    )
    add_figure(doc, "week4_03_procurement_vs_schedule_delay.png", "Figure 3. Procurement delay versus schedule delay; the line summarizes a linear association, not a causal effect.")
    doc.add_paragraph("The procurement view is actionable because order placement, submittal approval, promised delivery, required-on-site date, and alternative sourcing can often be managed before a late delivery appears as lost site time. An exception list should show the item, affected work package, days of slippage, schedule float remaining, owner, and next decision date.")
    add_figure(doc, "week4_04_budget_vs_actual_cost.png", "Figure 4. Planned budget versus actual final cost in the Week 2 synthetic sample. Points above the parity line exceed budget.")
    doc.add_paragraph(f"Mean cost overrun was {w2['mean_cost_overrun_pct']:.2f}% in the Week 2 sample, while {w2['over_10pct_cost_overrun_rate_pct']:.1f}% of records exceeded the 10% threshold. A cost-control meeting should reconcile commitments, actual costs, approved changes, pending claims, material-price exposure, forecast quantities, and contingency. The final cost chart alone is retrospective; early management requires a regularly refreshed estimate-at-completion.")

    # Models
    add_heading(doc, "4. Finding Three: Predictive Models Should Support, Not Replace, Review", 1)
    doc.add_paragraph("Week 3 compared two delay regressors and two classifiers for a cost-overrun flag of at least 10%. The metrics below come from the separate 500-row synthetic risk sample, with an 80/20 train-test split and five-fold cross-validation on the training sample. The regression target was delay in days; the classifier target was a thresholded cost-overrun flag. Recall measures how many overrun cases are flagged, while precision measures how many flagged cases are true; the alert threshold should reflect the cost of missed risks versus false alerts.")
    add_figure(doc, "week4_05_model_validation_tradeoffs.png", "Figure 5. Holdout model metrics; regression and classification scores have different meanings and scales.", width=5.4)
    add_heading(doc, "Delay forecast comparison", 2)
    add_table(doc, ["Model", "Holdout MAE (days)", "Holdout RMSE (days)", "Holdout R²", "Mean CV R² ± SD"], [
        ["Linear Regression", f"{w3['delay_linear_regression']['holdout_MAE_days']:.2f}", f"{w3['delay_linear_regression']['holdout_RMSE_days']:.2f}", f"{w3['delay_linear_regression']['holdout_R2']:.3f}", f"{w3['delay_linear_regression']['cv_R2_mean']:.3f} ± {w3['delay_linear_regression']['cv_R2_std']:.3f}"],
        ["Random Forest Regressor", f"{w3['delay_random_forest']['holdout_MAE_days']:.2f}", f"{w3['delay_random_forest']['holdout_RMSE_days']:.2f}", f"{w3['delay_random_forest']['holdout_R2']:.3f}", f"{w3['delay_random_forest']['cv_R2_mean']:.3f} ± {w3['delay_random_forest']['cv_R2_std']:.3f}"],
    ], widths=[1.55, 1.2, 1.2, 1.0, 1.55], font_size=8)
    doc.add_paragraph("In this example, Linear Regression had lower holdout error and higher holdout R² than the Random Forest Regressor. Its five-fold mean R² was also higher and less variable in the supplied runs. That makes the linear model a defensible baseline for this dataset—not a universal winner. Real projects can contain nonlinear relationships, reporting shifts, and differences between project types or phases.")
    add_heading(doc, "Cost-overrun risk classification", 2)
    add_table(doc, ["Model", "Accuracy", "Precision", "Recall", "F1", "ROC-AUC"], [
        ["Logistic Regression", f"{w3['cost_risk_logistic_regression']['accuracy']:.3f}", f"{w3['cost_risk_logistic_regression']['precision']:.3f}", f"{w3['cost_risk_logistic_regression']['recall']:.3f}", f"{w3['cost_risk_logistic_regression']['f1']:.3f}", f"{w3['cost_risk_logistic_regression']['roc_auc']:.3f}"],
        ["Random Forest Classifier", f"{w3['cost_risk_random_forest']['accuracy']:.3f}", f"{w3['cost_risk_random_forest']['precision']:.3f}", f"{w3['cost_risk_random_forest']['recall']:.3f}", f"{w3['cost_risk_random_forest']['f1']:.3f}", f"{w3['cost_risk_random_forest']['roc_auc']:.3f}"],
    ], widths=[2.0, 0.75, 0.75, 0.7, 0.65, 0.8], font_size=7.8)

    # Scenario and risk
    add_heading(doc, "5. Finding Four: What-if Scenarios Can Organize Mitigation Discussions", 1)
    doc.add_paragraph("Week 3 evaluated a single simulated high-risk project under several hypothetical input changes. The exercise illustrates how a model can help compare possible response plans before a meeting. It does not estimate causal intervention effects: changing a model input is not proof that a real intervention will produce the same result.")
    add_figure(doc, "week4_06_mitigation_scenario_comparison.png", "Figure 6. Predicted delay for one simulated high-risk project under selected what-if scenarios.")
    baseline = scenarios.loc[scenarios.scenario.eq("Baseline high-risk project")].iloc[0]
    combined = scenarios.loc[scenarios.scenario.eq("Combined controls")].iloc[0]
    delay_reduction = float(baseline.predicted_delay_days - combined.predicted_delay_days)
    pct_reduction = delay_reduction / float(baseline.predicted_delay_days) * 100
    add_table(doc, ["Scenario", "Predicted delay (days)", "Predicted cost-overrun probability", "Interpretation"], [
        ["Baseline high-risk project", f"{baseline.predicted_delay_days:.2f}", f"{100*baseline.predicted_cost_overrun_probability:.2f}%", "Reference case before hypothetical changes"],
        ["Supplier coordination", f"{float(scenarios.loc[scenarios.scenario.str.startswith('Supplier'), 'predicted_delay_days'].iloc[0]):.2f}", f"{100*float(scenarios.loc[scenarios.scenario.str.startswith('Supplier'), 'predicted_cost_overrun_probability'].iloc[0]):.2f}%", "Illustrative effect of reducing supplier-delay input"],
        ["Design freeze", f"{float(scenarios.loc[scenarios.scenario.str.startswith('Design'), 'predicted_delay_days'].iloc[0]):.2f}", f"{100*float(scenarios.loc[scenarios.scenario.str.startswith('Design'), 'predicted_cost_overrun_probability'].iloc[0]):.2f}%", "Illustrative effect of reducing design-change count"],
        ["Workforce plan", f"{float(scenarios.loc[scenarios.scenario.str.startswith('Workforce'), 'predicted_delay_days'].iloc[0]):.2f}", f"{100*float(scenarios.loc[scenarios.scenario.str.startswith('Workforce'), 'predicted_cost_overrun_probability'].iloc[0]):.2f}%", "Illustrative effect of higher labor availability"],
        ["Combined controls", f"{combined.predicted_delay_days:.2f}", f"{100*combined.predicted_cost_overrun_probability:.2f}%", f"{delay_reduction:.2f} fewer model-predicted days vs baseline ({pct_reduction:.1f}%)"],
    ], widths=[1.45, 1.05, 1.35, 2.1], font_size=7.7)

    add_heading(doc, "6. Risk Priorities and Management Response", 1)
    doc.add_paragraph("The matrix below is a structured prioritization starting point, using qualitative likelihood and impact ratings from 1 to 5. It is not a measured risk probability and should be discussed and adjusted with people who know the specific site, contract, schedule, and safety plan.")
    add_figure(doc, "week4_07_qualitative_risk_priority_matrix.png", "Figure 7. Qualitative likelihood-impact matrix. Rating positions are proposed for stakeholder review.")
    risk_rows = []
    for _, row in risks.iterrows():
        risk_rows.append([row.risk_area, f"{int(row.likelihood_1_to_5)} x {int(row.impact_1_to_5)} = {int(row.priority_score)}", row.priority_band, row.owner, row.first_response])
    add_table(doc, ["Risk area", "L x I", "Priority", "Accountable owner", "First response"], risk_rows, widths=[1.25, 0.65, 0.9, 1.25, 2.0], font_size=7.3)
    add_callout(doc, "Safety is different", "A risk matrix or predictive score must not be used to relax mandated safety controls. Safety reviews, permits, competent-person checks, incident investigation, and corrective actions remain mandatory regardless of the model output.", fill="FBEAEA")

    # Roadmap and conclusion
    add_heading(doc, "7. 30-60-90 Day Implementation Roadmap", 1)
    rec_rows = []
    for _, row in recs.iterrows():
        rec_rows.append([row.horizon, row.action, row.owner, row.measure_of_success])
    add_table(doc, ["Horizon", "Action", "Owner", "Evidence of progress"], rec_rows, widths=[0.85, 2.35, 1.25, 1.65], font_size=7.4)
    add_heading(doc, "Proposed weekly dashboard cadence", 2)
    for b in [
        "Before the meeting: data owners refresh actuals, commitments, forecasts, critical procurement items, open changes, labor plan, constraints, rework, and safety actions; include a timestamp and data-quality status.",
        "During the meeting: review exceptions first, assign owner and due date, record the decision, and note the expected effect on cost, finish date, quality, and safety.",
        "After the meeting: track open actions to closure and preserve forecast snapshots so prediction quality and management response can be evaluated later."
    ]: add_bullet(doc, b)
    add_heading(doc, "8. Explaining the Python workflow to non-technical users", 1)
    doc.add_paragraph("Python acts as a repeatable workbench. pandas reads the project tables and groups observations by type or risk condition. NumPy supports calculations used in the indicators and charts. Matplotlib turns the summaries into plots that help a meeting focus on differences and exceptions. scikit-learn prepares input columns, fits candidate models on a training subset, and evaluates them on held-out records. Cross-validation repeats the model check across several folds to show whether results are unusually dependent on one split.")
    doc.add_paragraph("The KPI dashboard is a descriptive snapshot; the scatterplots show associations; the regression model estimates a numeric delay; the classifier estimates the probability of a threshold-based risk label; and scenario analysis shows how predictions change when inputs are modified. These outputs answer different questions and should not be conflated. Model scores do not prove root cause, and a scenario result is not a guaranteed saving.")
    add_heading(doc, "9. Limitations and Governance", 1)
    for b in [
        "All datasets in this internship portfolio are simulated. There are no verified real construction-site observations in the report, so findings are demonstration results only.",
        "The Week 1, Week 2 and Week 3 datasets have different rows and generation assumptions. Their rates should not be combined or treated as a before/after series.",
        "Model metrics can be optimistic or unstable when features encode synthetic rules. Real deployment needs representative historic data, consistent definitions, quality controls and external validation.",
        "The scenario analysis is model-based what-if testing, not a causal experiment. Cost, feasibility, contractual terms, and unintended consequences must be reviewed by project experts.",
        "Real data access should be authorized; personally or commercially sensitive fields should be minimized; model changes, exclusions and decisions should be logged; and a named owner should monitor performance." 
    ]: add_bullet(doc, b)
    add_heading(doc, "Conclusion", 1)
    doc.add_paragraph("The four-week workflow demonstrates a practical path from data acquisition planning and quality checks to descriptive analysis, risk forecasting, visualization, and management action. Within the synthetic samples, procurement delay was positively associated with schedule delay, project-type delay varied, and simple baseline models provided useful examples of how delay and cost-overrun screening can be evaluated. The most important deliverable is the proposed operating discipline: maintain reliable data, discuss exceptions with accountable owners, measure forecast errors on real held-out projects before operational use, and keep safety and professional judgement central to every decision.")
    add_heading(doc, "Appendix A. Key metric definitions", 1)
    add_table(doc, ["Metric", "Definition", "How a team may use it"], [
        ["Schedule delay (days)", "Actual/forecast completion duration beyond the planned baseline in the source sample.", "Identify slippage and its effect on critical path and handover."],
        ["Cost overrun (%)", "Percentage difference between final cost and planned budget in the Week 2 synthetic sample.", "Compare current forecast with approved baseline and understand change exposure."],
        ["MAE", "Average absolute difference between predicted and observed values.", "Express typical delay forecast error in days."],
        ["RMSE", "Square root of average squared prediction error; larger errors receive more weight.", "Understand whether some large forecast misses are occurring."],
        ["R²", "Fraction of target variation explained by a fitted model relative to a mean baseline, under the evaluation setup.", "Compare baseline fit, not a guarantee of accuracy for an individual project."],
        ["Precision / recall", "Precision: share of predicted risk flags that are true; recall: share of actual risk cases detected.", "Balance false alarms against missed high-risk cases."],
        ["ROC-AUC", "Measures how well a classifier ranks positive cases above negative cases across thresholds.", "Compare ranking ability; it does not choose the operational threshold."],
    ], widths=[1.25, 2.65, 2.3], font_size=7.8)
    add_note(doc, "Prepared from the project files generated for Weeks 1-3. No external operational dataset was introduced into this final synthesis.")

    OUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUT)
    print(f"Wrote {OUT}")

if __name__ == "__main__":
    main()
